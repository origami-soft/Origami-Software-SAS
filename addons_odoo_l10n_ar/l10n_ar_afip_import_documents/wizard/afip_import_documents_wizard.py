# -*- encoding: utf-8 -*-

from odoo import models, fields
from odoo.exceptions import ValidationError
from datetime import datetime
from itertools import groupby
from dateutil import relativedelta
from io import BytesIO
import openpyxl, base64


class AfipImportDocumentsWizard(models.TransientModel):

    _name = 'afip.import.documents.wizard'

    file = fields.Binary(
        string='Archivo',
        help='Excel bajado de los comprobantes emitidos/recibidos de afip',
        required=True
    )
    filename = fields.Char(
        string='Nombre de archivo'
    )
    type = fields.Selection(
        string='Tipo',
        selection=[('sent', 'Emitidos'), ('received', 'Recibidos')],
        required=True
    )
    company_id = fields.Many2one(
        comodel_name='res.company',
        string='Empresa',
        required=True,
        default=lambda self: self.env.company
    )
    document_line_ids = fields.One2many(
        comodel_name='afip.import.documents.line.wizard',
        inverse_name='wizard_id',
        string='Comprobantes'
    )

    def validate_documents(self):
        self.document_line_ids.search([]).unlink()
        if self.type == 'sent':
            return self.validate_sent_documents()
        elif self.type == 'received':
            return self.validate_received_documents()

    def validate_sent_documents(self):
        vals = self.get_xls_values()
        non_found_documents = self.document_line_ids
        documents = self.document_line_ids = self.document_line_ids.create(vals)
        if documents:
            invoices = self._search_invoices()
            sorted_documents = documents.sorted(key=lambda x: (x.voucher_type, x.point_of_sale, x.voucher_name))
            grouped_documents = groupby(sorted_documents, key=lambda x: (x.voucher_type, x.point_of_sale))
            for (voucher_type, point_of_sale), grouped_document in grouped_documents:
                filtered_invoices = invoices.filtered(
                    lambda x: x.voucher_type_id.code == int(voucher_type.split(' ')[0]) and
                    x.pos_ar_id.name.lstrip("0") == point_of_sale.lstrip("0")
                )

                grouped_document = list(grouped_document)
                max_document = max(int(document.voucher_name) for document in grouped_document)
                min_document = min(int(document.voucher_name) for document in grouped_document)
                invoice_numbers = [
                    int(voucher_name.split('-')[1]) for voucher_name in filtered_invoices.mapped('voucher_name')
                ]
                non_found_numbers = [
                    number for number in range(min_document, max_document+1) if number not in invoice_numbers
                ]
                non_found_documents |= self.document_line_ids.filtered(
                    lambda x: x.voucher_type == voucher_type and x.point_of_sale == point_of_sale
                    and int(x.voucher_name) in non_found_numbers
                )

        self.document_line_ids = non_found_documents
        return self.get_create_invoice_view()

    def validate_received_documents(self):
        vals = self.get_xls_values()
        non_found_documents = self.document_line_ids
        documents = self.document_line_ids = self.document_line_ids.create(vals)
        if documents:
            invoices = self._search_invoices()
            for document in documents:
                invoice = invoices.filtered(
                    lambda x: document.document_number == x.partner_id.vat and
                    document.point_of_sale.lstrip("0") == x.voucher_name.split('-')[0].lstrip("0") and
                    document.voucher_name.lstrip("0") == x.voucher_name.split('-')[1].lstrip("0")
                )
                if not invoice:
                    non_found_documents |= document

        self.document_line_ids = non_found_documents
        return self.get_create_invoice_view()

    def create_invoices(self):
        invoices = []

        # Para optimizar, nos armamos listas valores únicos con lo que hay que mapear con odoo
        voucher_types = list(set(
            voucher_type.split(' ')[0] for voucher_type in set(self.document_line_ids.mapped('voucher_type'))
        ))
        vat_numbers = list(set(self.document_line_ids.mapped('document_number')))
        currencies = list(set(self.document_line_ids.mapped('currency')))
        partners = self.env['res.partner'].search([('vat', 'in', vat_numbers)])
        voucher_types = self.env['voucher.type'].search([('code', 'in', voucher_types)])

        if self.type == 'sent':
            documents_pos = self.document_line_ids.mapped('point_of_sale')
            journals_with_pos = self.env['account.journal'].search([
                ('pos_ar_id', '!=', False),
                ('company_id', '=', self.company_id.id),
                ('type', '=', 'sale')
            ])
            point_of_sales_name = [name.lstrip("0") for name in journals_with_pos.mapped('pos_ar_id').mapped('name')]
            for document_pos in documents_pos:
                if document_pos.lstrip("0") not in point_of_sales_name:
                    raise ValidationError("No se encontro punto de venta {} en odoo".format(document_pos))

        non_found_partners = [partner for partner in vat_numbers if partner not in list(set(partners.mapped('vat')))]

        if non_found_partners:
            raise ValidationError("No se encontraron partners para los numeros de documentos:\n {}".format(
                '\n'.join(non_found_partners))
            )

        odoo_currencies = {}
        for currency in currencies:
            # TODO: Ver como solucionamos esto, en el excel vienen distintos codigos que lo informado en afip
            if currency == '$':
                odoo_currency = 'PES'
            elif currency == 'USD':
                odoo_currency = 'DOL'
            else:
                odoo_currency = currency

            try:
                odoo_currency = self.env['codes.models.relation'].get_record_from_code(
                    'res.currency',
                    odoo_currency,
                    'Afip'
                )
            except Exception:
                raise ValidationError('No se encontró moneda con código {}'.format(currency))
            odoo_currencies[currency] = odoo_currency

        for document in self.document_line_ids:
            lines = document.get_invoice_lines()
            voucher_type = int(document.voucher_type.split(' ')[0])
            invoice_voucher_type = voucher_types.filtered(lambda x: x.code == voucher_type)
            partner = partners.filtered(lambda x: x.vat == document.document_number)[0]
            vals = {
                'voucher_type_id': invoice_voucher_type.id,
                'partner_id': partner.id,
                'fiscal_position_id': partner.property_account_position_id.id,
                'invoice_date': document.date,
                'date': document.date,
                'currency_id': odoo_currencies[document.currency].id,
                'currency_rate': document.currency_value,
                'invoice_line_ids': lines
            }
            if self.type == 'sent':
                journal = journals_with_pos.filtered(
                    lambda x: x.pos_ar_id.name.lstrip("0") == document.point_of_sale.lstrip("0")
                )[0]
                document_book = journal.pos_ar_id.document_book_ids.filtered(
                    lambda x: x.voucher_type_id == invoice_voucher_type
                )
                vals.update({
                    'voucher_name': "{:0>{prefix_qty}}-{:0>8}".format(
                        document.point_of_sale,
                        document.voucher_name,
                        prefix_qty=journal.pos_ar_id.prefix_quantity or 0,
                    ),
                    'journal_id': journal.id,
                    'move_type': 'out_refund' if invoice_voucher_type.category == 'refund' else 'out_invoice',
                    'cae': document.cae,
                    'cae_due_date': document.date + relativedelta.relativedelta(days=10),
                    'document_book_id': document_book and document_book[0].id
                })
            else:
                vals.update({
                    'voucher_name': "{}-{}".format(document.point_of_sale, document.voucher_name),
                    'move_type': 'in_refund' if invoice_voucher_type.category == 'refund' else 'in_invoice'
                })
            invoices.append(vals)

        if invoices:
            invoices = self.env['account.move'].create(invoices)
            for invoice in invoices.filtered(lambda x: x.move_type in ['out_invoice', 'out_refund']):
                invoice.set_voucher_name()
            return {
                'type': 'ir.actions.act_window',
                'name': 'Factura creadas',
                'res_model': 'account.move',
                'views': [[self.env.ref('account.view_invoice_tree').id, "tree"], [False, "form"]],
                'domain': [('id', 'in', invoices.ids)],
                'context': {'default_type': 'in_invoice'}
            }

    def get_xls_values(self):
        try:
            book = openpyxl.load_workbook(BytesIO(base64.b64decode(self.file)))
            sheet = book.worksheets[0]
            vals = []
            for x in range(3, sheet.max_row + 1):
                vals.append({
                    'date': datetime.strptime(sheet.cell(x, 1).value, '%d/%m/%Y'),
                    'voucher_type': sheet.cell(x, 2).value,
                    'point_of_sale': str(int(sheet.cell(x, 3).value)),
                    'voucher_name': str(int(sheet.cell(x, 4).value)),
                    'cae': str(int(sheet.cell(x, 6).value)),
                    'document_type': sheet.cell(x, 7).value,
                    'document_number': str(int(sheet.cell(x, 8).value)),
                    'name': sheet.cell(x, 9).value,
                    'currency_value': sheet.cell(x, 10).value,
                    'currency': sheet.cell(x, 11).value,
                    'amount_untaxed': sheet.cell(x, 12).value,
                    'amount_not_taxed': sheet.cell(x, 13).value,
                    'amount_exempt': sheet.cell(x, 14).value,
                    'amount_other_tributes': sheet.cell(x, 15).value,
                    'amount_vat': sheet.cell(x, 16).value,
                    'amount_total': sheet.cell(x, 17).value,
                })
        except Exception:
            raise ValidationError("Hubo un error al intentar leer el archivo.")
        return vals

    def get_create_invoice_view(self):
        view_id = self.env.ref('l10n_ar_afip_import_documents.afip_non_found_documents_wizard').id
        return {
            'type': 'ir.actions.act_window',
            'name': 'Documentos no encontrados',
            'view_mode': 'form',
            'res_model': self._name,
            'views': [[view_id, "form"]],
            'res_id': self.id,
            'target': 'new'
        }

    def _search_invoices(self):
        sorted_documents = self.document_line_ids.sorted(key=lambda x: x.date)
        date_from = sorted_documents[0].date
        date_to = sorted_documents[-1:].date
        return self.env['account.move'].search([
            ('company_id', '=', self.company_id.id),
            ('date', '>=', date_from),
            ('date', '<=', date_to),
            ('move_type', 'in', ['in_invoice', 'in_refund'] if self.type == 'received' else ['out_invoice', 'out_refund']),
            ('voucher_name', 'like', '-'),
        ])

# vim:expandtab:smartindent:tabstop=4:softtabstop=4:shiftwidth=4:

